"""Render policy versions to Markdown, HTML, DOCX and PDF with byte-stable output."""

from __future__ import annotations

import html
import io
import zipfile
from datetime import UTC, datetime

from data.generator.catalog import DOMAIN_TITLES, TENANT
from data.generator.corpus import Policy, Version

FIXED_TS = datetime(2026, 1, 1, tzinfo=UTC)
EXT = {"md": "md", "html": "html", "docx": "docx", "pdf": "pdf"}


def _meta_line(policy: Policy, version: Version) -> str:
    return (
        f"{TENANT} · {DOMAIN_TITLES[policy.domain]} policy · Document {policy.document_id} · "
        f"Version {version.version} · Owner: {policy.owner}"
    )


def _article_heading(number: int, heading: str, dash: str = "—") -> str:
    return f"Article {number} {dash} {heading}"


def render_md(policy: Policy, version: Version) -> bytes:
    lines = [f"# {policy.title}", "", _meta_line(policy, version), ""]
    for article in version.articles:
        lines += [f"## {_article_heading(article.number, article.heading)}", ""]
        for paragraph in article.paragraphs:
            lines += [paragraph, ""]
    return "\n".join(lines).encode()


def render_html(policy: Policy, version: Version) -> bytes:
    esc = html.escape
    parts = [
        "<!doctype html>",
        '<html lang="en"><head><meta charset="utf-8">',
        f"<title>{esc(policy.title)}</title></head><body>",
        f"<h1>{esc(policy.title)}</h1>",
        f'<p class="meta">{esc(_meta_line(policy, version))}</p>',
    ]
    for article in version.articles:
        parts.append(f"<h2>{esc(_article_heading(article.number, article.heading))}</h2>")
        parts.extend(f"<p>{esc(p)}</p>" for p in article.paragraphs)
    parts.append("</body></html>")
    return "\n".join(parts).encode()


def _stable_zip(blob: bytes) -> bytes:
    """Rewrite a zip with fixed timestamps and sorted entries so DOCX bytes are reproducible."""
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        names = src.namelist()
        ordered = ["[Content_Types].xml"] + sorted(n for n in names if n != "[Content_Types].xml")
        for name in ordered:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, src.read(name))
    return out.getvalue()


def render_docx(policy: Policy, version: Version) -> bytes:
    from docx import Document

    doc = Document()
    props = doc.core_properties
    props.created = FIXED_TS
    props.modified = FIXED_TS
    props.last_printed = FIXED_TS
    props.author = TENANT
    props.last_modified_by = TENANT
    props.revision = 1
    props.title = policy.title
    doc.add_heading(policy.title, level=1)
    doc.add_paragraph(_meta_line(policy, version))
    for article in version.articles:
        doc.add_heading(_article_heading(article.number, article.heading), level=2)
        for paragraph in article.paragraphs:
            doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return _stable_zip(buf.getvalue())


def _latin1(text: str) -> str:
    text = text.replace("—", "-").replace("​", "").replace("’", "'")
    return text.encode("latin-1", "replace").decode("latin-1")


def render_pdf(policy: Policy, version: Version) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_creation_date(FIXED_TS)
    pdf.set_title(_latin1(policy.title))
    pdf.set_author(TENANT)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    width = pdf.epw
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(width, 7, _latin1(policy.title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.multi_cell(width, 5, _latin1(_meta_line(policy, version)), new_x="LMARGIN", new_y="NEXT")
    for article in version.articles:
        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(
            width,
            6,
            _latin1(_article_heading(article.number, article.heading, "-")),
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.set_font("Helvetica", "", 10)
        for paragraph in article.paragraphs:
            pdf.multi_cell(width, 5, _latin1(paragraph), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
    return bytes(pdf.output())


RENDERERS = {"md": render_md, "html": render_html, "docx": render_docx, "pdf": render_pdf}


def render_version(policy: Policy, version: Version, fmt: str | None = None) -> bytes:
    return RENDERERS[fmt or version.source_format](policy, version)


def source_path(policy: Policy, version: Version) -> str:
    return f"raw/{policy.document_id}/v{version.version}.{EXT[version.source_format]}"


def render_policy(policy: Policy) -> dict[str, bytes]:
    return {source_path(policy, v): render_version(policy, v) for v in policy.versions}
